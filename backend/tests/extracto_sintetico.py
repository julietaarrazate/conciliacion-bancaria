"""Archivos de extracto SINTÉTICOS para los tests del lector Excel/CSV.

Reproducen la estructura de extractos reales (título arriba, encabezado en la
fila 3, columna Orden con fórmulas y valor calculado guardado, del más nuevo al
más viejo, fila TOTAL al final, columnas extra) con datos inventados. Nunca
copiar datos de un extracto real acá.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import date, timedelta
from decimal import Decimal

import openpyxl

SALDO_INICIAL = Decimal("1000.00")


def movimientos_cronologicos(n: int = 6) -> list[dict]:
    """Movimientos inventados del más viejo al más nuevo, con saldo corrido."""
    importes = [Decimal("1500.50"), Decimal("-200.00"), Decimal("3000.00"),
                Decimal("-45.25"), Decimal("780.10"), Decimal("-1234.56"),
                Decimal("99.99"), Decimal("-10.00")]
    saldo = SALDO_INICIAL
    movs = []
    for i in range(n):
        monto = importes[i % len(importes)]
        saldo += monto
        movs.append({
            "orden": 100 + i,
            "fecha": date(2026, 5, 4) + timedelta(days=i),
            "titular": f"CLIENTE PRUEBA {i} 2012345678{i}",
            "monto": monto,
            "saldo": saldo,
        })
    return movs


def _con_formulas_en_orden(contenido: bytes, filas: list[int]) -> bytes:
    """Convierte las celdas A<fila> en fórmulas `=A<fila+1>+1` conservando el valor
    calculado (`<v>`), como queda un .xlsx guardado por Excel."""
    entrada = zipfile.ZipFile(io.BytesIO(contenido))
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as z:
        for item in entrada.infolist():
            data = entrada.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                for f in filas:
                    xml = re.sub(
                        rf'(<c r="A{f}"[^>]*>)(<v>)',
                        rf"\g<1><f>A{f + 1}+1</f>\g<2>",
                        xml,
                    )
                data = xml.encode("utf-8")
            z.writestr(item, data)
    return salida.getvalue()


def xlsx_estilo_macro(n: int = 6, *, titulo: str = "Extracto Banco Macro Mayo.xlsx",
                      romper_fila: int | None = None) -> bytes:
    """Hoja "Extracto conciliado": título, fila vacía, encabezado en la fila 3
    (Orden, Fecha, Mes, Titular, Importe, Saldo + extras), del más nuevo al más
    viejo, Orden con fórmulas y una fila TOTAL al final.
    `romper_fila`: índice (en el orden del archivo) cuyo saldo se altera."""
    movs = list(reversed(movimientos_cronologicos(n)))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Extracto conciliado"
    ws.cell(1, 1).value = titulo
    for c, h in enumerate(["Orden", "Fecha", "Mes", "Titular", "Importe", "Saldo",
                           "Cliente acreditado", "Fecha acred."], start=1):
        ws.cell(3, c).value = h
    for i, m in enumerate(movs):
        r = 4 + i
        saldo = m["saldo"] + (Decimal("50") if romper_fila == i else 0)
        ws.cell(r, 1).value = m["orden"]
        ws.cell(r, 2).value = m["fecha"]
        ws.cell(r, 3).value = m["fecha"].month
        ws.cell(r, 4).value = m["titular"]
        ws.cell(r, 5).value = float(m["monto"])
        ws.cell(r, 6).value = float(saldo)
    total = 4 + len(movs)
    ws.cell(total, 4).value = "TOTAL"
    ws.cell(total, 5).value = float(sum(m["monto"] for m in movs))
    ws.cell(total, 7).value = "nota"
    buf = io.BytesIO()
    wb.save(buf)
    # Orden: todas menos la última fila de datos son fórmulas (=A<r+1>+1).
    return _con_formulas_en_orden(buf.getvalue(), list(range(4, total - 1)))


def xlsx_desde_filas(filas: list[list], titulo_hoja: str = "Hoja1") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = titulo_hoja
    for fila in filas:
        ws.append(fila)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
