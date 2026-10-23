"""Entradas de terceros (revisión de seguridad oct/2026).

- Inyección de fórmulas: un titular/concepto que empieza con `=`, `+`, `-` o `@` se
  exporta como texto (con apóstrofo), no como fórmula viva. Números/fechas no cambian.
- Subidas con tope de tamaño (413 sin leer el archivo entero) y chequeo de zip bomb en
  .xlsx antes de pasarlo a openpyxl/pandas.
- PDF: un nombre con `&` o `<b>` no rompe reportlab (se escapa el dato, no el markup propio).
- DELETE /push/subscribe borra solo la suscripción del usuario actual, no la de otro.
"""
import csv
import io
import zipfile
from datetime import date
from decimal import Decimal

import openpyxl
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.main import app
from app.models.organizacion import Organizacion
from app.models.push_subscription import PushSubscription
from app.models.user import User
from app.routers import planillas as planillas_router
from app.services.auth import create_access_token, get_password_hash
from app.services.excel_export import export_movimientos
from app.services.excel_safe import safe_cell, sanear_workbook, verificar_xlsx
from app.services.export_contable import _export_csv, _export_regisoft
from app.services.limites import MAX_BYTES_PLANILLA
from app.services.pdf_export import (
    cierre_mensual_pdf, cuenta_corriente_pdf, estado_cuenta_pdf, export_planilla_conciliada_pdf,
    recibo_sueldo_pdf,
)


@pytest.fixture
def db():
    engine = create_engine(
        "sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    s = sessionmaker(bind=engine)()
    s.add_all([Organizacion(id=1, nombre="Org1"), Organizacion(id=2, nombre="Org2")])
    s.commit()
    for email, org in (("op1@t.test", 1), ("op2@t.test", 2)):
        s.add(User(
            email=email, full_name=email, hashed_password=get_password_hash("pw123456"),
            organizacion_id=org, is_active=True, role="operador", is_superadmin=False,
        ))
    s.commit()
    yield s
    s.close()


@pytest.fixture
def client(db):
    app.dependency_overrides[get_db] = lambda: db
    planillas_router.limiter.reset()
    yield TestClient(app, raise_server_exceptions=True)
    app.dependency_overrides.pop(get_db, None)


def _auth(db, email):
    u = db.query(User).filter(User.email == email).first()
    tok = create_access_token({"sub": u.email, "user_id": u.id, "role": u.role})
    return {"Authorization": f"Bearer {tok}"}


# ── 1. Inyección de fórmulas en Excel/CSV ────────────────────────────────────

def test_safe_cell_neutraliza_prefijos_de_formula():
    for v in ("=HYPERLINK(\"http://x\",\"y\")", "+1+cmd|' /C calc'!A0", "-2+3", "@SUM(A1)",
              "\t=1", "\r=1"):
        assert safe_cell(v) == "'" + v


def test_safe_cell_no_toca_numeros_fechas_ni_texto_comun():
    assert safe_cell(Decimal("-12.50")) == Decimal("-12.50")
    assert safe_cell(-3) == -3
    assert safe_cell(1.5) == 1.5
    assert safe_cell(date(2026, 1, 2)) == date(2026, 1, 2)
    assert safe_cell(None) is None
    assert safe_cell("") == ""
    assert safe_cell("Juan Pérez") == "Juan Pérez"
    # Texto que es solo un número (o el guion de "sin dato") no puede armar una fórmula.
    assert safe_cell("-1234,56") == "-1234,56"
    assert safe_cell("-") == "-"


def test_sanear_workbook_deja_formulas_como_texto():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "=1+1"
    ws["B1"] = Decimal("10.00")
    sanear_workbook(wb)
    assert ws["A1"].value == "'=1+1"
    assert ws["A1"].data_type == "s"
    assert ws["B1"].value == Decimal("10.00")


def test_export_movimientos_titular_malicioso_no_es_formula():
    xlsx = export_movimientos("ext", [{
        "orden": 1, "fecha": date(2026, 1, 5), "mes": 1,
        "titular": '=HYPERLINK("http://evil","clic")', "monto": Decimal("-100.50"), "saldo": Decimal("5"),
        "cliente_acreditado": "@cliente",
    }])
    ws = openpyxl.load_workbook(io.BytesIO(xlsx)).active
    assert ws.cell(6, 4).value == "'=HYPERLINK(\"http://evil\",\"clic\")"
    assert ws.cell(6, 4).data_type == "s"
    assert ws.cell(6, 7).value == "'@cliente"
    # El importe negativo sigue siendo número con su formato.
    assert ws.cell(6, 5).value == -100.5
    assert ws.cell(6, 5).number_format == '"$"#,##0.00'


def _fila_contable(desc):
    return {
        "fecha": date(2026, 3, 1), "numero_asiento": 7, "codigo_cuenta": "1.1.01",
        "nombre_cuenta": "=cmd", "debe": Decimal("10.00"), "haber": Decimal("0"), "descripcion": desc,
    }


def test_export_contable_csv_neutraliza_descripcion():
    content, _, _ = _export_csv([_fila_contable("=HYPERLINK(1)")], ";", "utf-8", "DD/MM/YYYY", {})
    filas = list(csv.reader(io.StringIO(content.decode("utf-8")), delimiter=";"))
    assert filas[1][3] == "'=cmd"
    assert filas[1][6] == "'=HYPERLINK(1)"
    assert filas[1][4] == "10.00"


def test_export_regisoft_neutraliza_descripcion():
    content, _, _ = _export_regisoft([_fila_contable("+SUM(A1)")])
    filas = list(csv.reader(io.StringIO(content.decode("utf-8"))))
    assert filas[1][2] == "'+SUM(A1)"


# ── 2. Tope de tamaño y zip bomb en subidas ──────────────────────────────────

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx_real() -> bytes:
    wb = openpyxl.Workbook()
    wb.active.append(["Fecha", "Titular", "Monto"])
    wb.active.append(["01/01/2026", "Juan", 100])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _zip_bomb(mb: int) -> bytes:
    """ZIP con un único archivo interno de `mb` MB de ceros (comprime ~1000:1)."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        with zf.open("xl/worksheets/sheet1.xml", "w") as dst:
            bloque = b"0" * (1024 * 1024)
            for _ in range(mb):
                dst.write(bloque)
    return buf.getvalue()


def test_verificar_xlsx_acepta_un_excel_normal_y_lo_que_no_es_zip():
    verificar_xlsx(_xlsx_real())
    verificar_xlsx(b"fecha;monto\n01/01/2026;100\n")
    verificar_xlsx(b"PK\x03\x04 roto")


def test_verificar_xlsx_rechaza_tasa_de_compresion_anomala():
    with pytest.raises(HTTPException) as ex:
        verificar_xlsx(_zip_bomb(5))
    assert ex.value.status_code == 413


def test_verificar_xlsx_rechaza_descomprimido_enorme():
    bomba = _zip_bomb(101)
    assert len(bomba) < 1024 * 1024  # el archivo subido es chico
    with pytest.raises(HTTPException) as ex:
        verificar_xlsx(bomba)
    assert ex.value.status_code == 413
    assert "descomprimido" in ex.value.detail


def test_preview_planilla_rechaza_archivo_demasiado_grande(client, db):
    grande = b"a" * (MAX_BYTES_PLANILLA + 1)
    r = client.post("/planillas/preview", headers=_auth(db, "op1@t.test"),
                    files={"file": ("p.csv", grande, "text/csv")})
    assert r.status_code == 413
    assert "MB" in r.json()["detail"]


def test_preview_planilla_rechaza_zip_bomb(client, db):
    r = client.post("/planillas/preview", headers=_auth(db, "op1@t.test"),
                    files={"file": ("p.xlsx", _zip_bomb(5), _XLSX)})
    assert r.status_code == 413


@pytest.mark.parametrize("url", [
    "/iva/comprobantes/comparar?direccion=recibido&periodo=2026-01",
    "/cheques/importar",
    "/tarjetas/upload?marca=visa",
    "/extractos/upload",
])
def test_otras_subidas_de_planillas_rechazan_zip_bomb(client, db, url):
    r = client.post(url, headers=_auth(db, "op1@t.test"),
                    files={"file": ("x.xlsx", _zip_bomb(5), _XLSX)})
    assert r.status_code == 413, r.text


def test_mis_comprobantes_de_varias_orgs_rechaza_zip_bomb(client, db):
    u = db.query(User).filter(User.email == "op1@t.test").first()
    u.role = "contador"
    db.commit()
    r = client.post("/procesos-masivos/mis-comprobantes/comparar?direccion=recibido&periodo=2026-01",
                    headers=_auth(db, "op1@t.test"),
                    files=[("files", ("x.xlsx", _zip_bomb(5), _XLSX))])
    assert r.status_code == 413, r.text


# ── 3. Escape de texto en los PDF (reportlab interpreta markup) ──────────────

_MALO = "A & B <b>x"


def _es_pdf(b: bytes) -> bool:
    return b.startswith(b"%PDF")


def test_pdf_planilla_conciliada_con_nombre_con_markup():
    pdf = export_planilla_conciliada_pdf({
        "cliente_nombre": _MALO, "nombre_archivo": "pl <i>.xlsx",
        "rows": [{"monto": Decimal("10.00"), "titular": _MALO, "status": "ok", "cuit": "20123456789"}],
        "total_declarado": None,
    })
    assert _es_pdf(pdf)


def test_pdf_estado_cuenta_con_nombre_con_markup():
    pdf = estado_cuenta_pdf({
        "cliente": {"nombre": _MALO, "cuit": "20-1 & <x>"},
        "periodo": {"desde": "2026-01-01", "hasta": "2026-01-31"},
        "resumen": {"conciliado_periodo": 10, "pendiente_periodo": 0},
        "planillas": [{"nombre_archivo": _MALO, "filas": []}],
        "cheques": [{"banco_origen": _MALO, "monto": 5}],
        "pagos": [{"concepto": _MALO, "monto": 1}],
    })
    assert _es_pdf(pdf)


def test_pdf_cierre_cuenta_corriente_y_recibo_con_markup():
    assert _es_pdf(cierre_mensual_pdf({
        "periodo_actual": {}, "top_clientes": [{"nombre": _MALO, "total": 1}],
        "cheques": {"proximos_a_vencer": [{"cliente_nombre": _MALO, "monto": 1}]},
    }, 2026, 1, org_nombre=_MALO))
    assert _es_pdf(cuenta_corriente_pdf({
        "cliente": {"nombre": _MALO}, "cuenta": {"codigo": "1.1", "nombre": _MALO},
        "movimientos": [{"fecha": "2026-01-02", "tipo_label": _MALO, "referencia": _MALO, "debito": 1}],
    }))
    assert _es_pdf(recibo_sueldo_pdf(
        {"sueldo_bruto": 100, "sueldo_neto": 80, "total_aportes": 20},
        {"nombre": _MALO, "cuil": "20-1", "categoria_nombre": _MALO, "convenio_nombre": "<x"},
        "2026-01",
    ))


# ── 5. DELETE /push/subscribe solo borra la suscripción propia ───────────────

def test_desuscribir_push_no_borra_la_suscripcion_de_otro_usuario(client, db):
    otro = db.query(User).filter(User.email == "op2@t.test").first()
    db.add(PushSubscription(user_id=otro.id, endpoint="https://push.test/ajeno", p256dh="k", auth="a"))
    db.commit()
    body = {"endpoint": "https://push.test/ajeno", "keys": {}}
    r = client.request("DELETE", "/push/subscribe", headers=_auth(db, "op1@t.test"), json=body)
    assert r.status_code == 200
    assert db.query(PushSubscription).filter_by(endpoint="https://push.test/ajeno").count() == 1

    # El dueño sí la puede borrar.
    r = client.request("DELETE", "/push/subscribe", headers=_auth(db, "op2@t.test"), json=body)
    assert r.status_code == 200
    assert db.query(PushSubscription).filter_by(endpoint="https://push.test/ajeno").count() == 0
